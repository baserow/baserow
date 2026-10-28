import codecs
import re
import zipfile
from email import policy
from email.parser import BytesParser
from html.parser import HTMLParser
from typing import IO, Iterable
from xml.etree.ElementTree import Element

from defusedxml.ElementTree import iterparse
from pypdf import PdfReader

from .registries import FileTextExtractorType

# Decoding can't produce more characters than bytes, so reading a few bytes per
# requested character is enough, even for multibyte encodings.
MAX_BYTES_PER_CHARACTER = 4

# A small compressed file can expand enormously, which takes too long to read.
MAX_UNCOMPRESSED_ARCHIVE_BYTES = 200 * 1024 * 1024

WORD_NAMESPACE = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
DRAWING_NAMESPACE = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
OPEN_DOCUMENT_TEXT_NAMESPACE = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
OPEN_DOCUMENT_TABLE_NAMESPACE = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"
LEGACY_OFFICE_FILE_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _read_text(file: IO[bytes], max_characters: int) -> str:
    content = file.read(max_characters * MAX_BYTES_PER_CHARACTER)
    # Incremental, because the content can end in the middle of a character.
    if content.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        text = codecs.getincrementaldecoder("utf-16")().decode(content)
    else:
        try:
            text = codecs.getincrementaldecoder("utf-8-sig")().decode(content)
        except UnicodeDecodeError:
            # The default encoding of Windows, of which the files are the most
            # common ones that aren't UTF-8.
            text = content.decode("cp1252", errors="replace")
    return text[:max_characters]


class PlainTextFileTextExtractorType(FileTextExtractorType):
    type = "plain_text"

    mime_types = {
        "application/json",
        "application/xml",
        "application/x-yaml",
        "application/yaml",
        "application/javascript",
        "application/sql",
    }
    extensions = (
        ".txt",
        ".md",
        ".markdown",
        ".csv",
        ".tsv",
        ".json",
        ".xml",
        ".yaml",
        ".yml",
        ".log",
        ".rst",
        ".ini",
        ".toml",
    )

    def can_extract(self, mime_type: str, file_name: str) -> bool:
        if mime_type == "text/html" or file_name.endswith((".html", ".htm")):
            return False
        return (
            mime_type.startswith("text/")
            or mime_type in self.mime_types
            or file_name.endswith(self.extensions)
        )

    def extract(self, file: IO[bytes], max_characters: int) -> str:
        return _read_text(file, max_characters)


class _HTMLTextCollector(HTMLParser):
    ignored_tags = {"script", "style", "head"}

    def __init__(self):
        super().__init__()
        self.parts = []
        self.ignored_depth = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.ignored_tags:
            self.ignored_depth += 1

    def handle_endtag(self, tag):
        if tag in self.ignored_tags and self.ignored_depth:
            self.ignored_depth -= 1

    def handle_data(self, data):
        if not self.ignored_depth and data.strip():
            self.parts.append(data.strip())


def _html_to_text(html: str, max_characters: int) -> str:
    collector = _HTMLTextCollector()
    collector.feed(html)
    collector.close()
    return "\n".join(collector.parts)[:max_characters]


class HTMLFileTextExtractorType(FileTextExtractorType):
    type = "html"

    def can_extract(self, mime_type: str, file_name: str) -> bool:
        return mime_type == "text/html" or file_name.endswith((".html", ".htm"))

    def extract(self, file: IO[bytes], max_characters: int) -> str:
        # Markup takes up a large part of an HTML file, so more of it is read to
        # still end up with roughly the requested number of text characters.
        return _html_to_text(_read_text(file, max_characters * 4), max_characters)


class PDFFileTextExtractorType(FileTextExtractorType):
    type = "pdf"

    def can_extract(self, mime_type: str, file_name: str) -> bool:
        return mime_type == "application/pdf" or file_name.endswith(".pdf")

    def extract(self, file: IO[bytes], max_characters: int) -> str:
        parts = []
        length = 0
        for page in PdfReader(file).pages:
            if length >= max_characters:
                break
            text = page.extract_text() or ""
            parts.append(text)
            length += len(text)
        return "\n\n".join(parts)[:max_characters]


def _join_until(parts: Iterable[str], separator: str, max_characters: int) -> str:
    """
    Joins the parts until the text is long enough, so that the remaining parts
    don't have to be read.

    :param parts: The parts of the text.
    :param separator: The text between the parts.
    :param max_characters: The maximum length of the text.
    :return: The joined text.
    """

    collected = []
    length = 0
    for part in parts:
        if length >= max_characters:
            break
        collected.append(part)
        length += len(part) + len(separator)
    return separator.join(collected)[:max_characters]


def _open_archive(file: IO[bytes]) -> zipfile.ZipFile | None:
    """
    :param file: The opened zip file.
    :return: The opened archive, or None if it expands to more than can be read.
    :raises zipfile.BadZipFile: If the file isn't a zip file.
    """

    archive = zipfile.ZipFile(file)
    # The declared sizes can be trusted, because reading stops at them.
    if (
        sum(member.file_size for member in archive.infolist())
        > MAX_UNCOMPRESSED_ARCHIVE_BYTES
    ):
        archive.close()
        return None
    return archive


def _iterparse_releasing(content: IO[bytes]) -> Iterable[tuple[str, Element]]:
    """
    Yields the start and end events of the elements, and removes every element from
    the tree after its end event, so that the memory doesn't grow with the number of
    elements. The tail of an element is available until the next event.

    :param content: The XML document.
    :return: The events and their elements.
    :raises ParseError: If the document isn't valid XML.
    """

    open_elements = []
    for event, element in iterparse(content, events=("start", "end")):
        if event == "start":
            yield event, element
            open_elements.append(element)
        else:
            open_elements.pop()
            yield event, element
            if open_elements:
                open_elements[-1].remove(element)


def _office_open_xml_paragraphs(document: IO[bytes], namespace: str) -> Iterable[str]:
    """
    Yields the paragraphs of a Word document or PowerPoint slide, which both store
    text in `p` paragraphs with `t` text elements.

    :param document: The XML of the document or slide.
    :param namespace: The XML namespace of the `p` and `t` elements.
    :return: The text of every paragraph that isn't empty.
    """

    # The text parts of the paragraphs that are being parsed, of which the last one
    # is nested in the others, like the paragraph of a text box.
    paragraphs = []
    # Streamed, so that a large document isn't decompressed further than needed.
    # Only `t` elements are visible text, unlike deleted text and field codes.
    for event, element in _iterparse_releasing(document):
        if element.tag == f"{namespace}p":
            if event == "start":
                paragraphs.append([])
            elif (text := "".join(paragraphs.pop())).strip():
                yield text
        elif element.tag == f"{namespace}t" and event == "end" and paragraphs:
            paragraphs[-1].append(element.text or "")


class WordFileTextExtractorType(FileTextExtractorType):
    type = "word"

    mime_types = {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.template",
        "application/vnd.ms-word.document.macroenabled.12",
    }

    def can_extract(self, mime_type: str, file_name: str) -> bool:
        return mime_type in self.mime_types or file_name.endswith(
            (".docx", ".docm", ".dotx")
        )

    def extract(self, file: IO[bytes], max_characters: int) -> str:
        archive = _open_archive(file)
        if archive is None:
            return ""
        with archive, archive.open("word/document.xml") as document:
            return _join_until(
                _office_open_xml_paragraphs(document, WORD_NAMESPACE),
                "\n",
                max_characters,
            )


class PowerPointFileTextExtractorType(FileTextExtractorType):
    type = "powerpoint"

    slide_name_pattern = re.compile(r"ppt/slides/slide(\d+)\.xml")

    def can_extract(self, mime_type: str, file_name: str) -> bool:
        return (
            mime_type
            == "application/vnd.openxmlformats-officedocument."
            "presentationml.presentation"
            or file_name.endswith(".pptx")
        )

    def extract(self, file: IO[bytes], max_characters: int) -> str:
        archive = _open_archive(file)
        if archive is None:
            return ""
        with archive:
            return _join_until(self._paragraphs(archive), "\n", max_characters)

    def _paragraphs(self, archive: zipfile.ZipFile) -> Iterable[str]:
        # The numbers of the slide files are roughly the slide order. The exact
        # order doesn't matter for searching, so the presentation isn't parsed.
        slides = sorted(
            (int(match.group(1)), name)
            for name in archive.namelist()
            if (match := self.slide_name_pattern.fullmatch(name))
        )
        for _, name in slides:
            with archive.open(name) as slide:
                yield from _office_open_xml_paragraphs(slide, DRAWING_NAMESPACE)


class OpenDocumentFileTextExtractorType(FileTextExtractorType):
    type = "open_document"

    mime_types = {
        "application/vnd.oasis.opendocument.text",
        "application/vnd.oasis.opendocument.spreadsheet",
        "application/vnd.oasis.opendocument.presentation",
    }

    def can_extract(self, mime_type: str, file_name: str) -> bool:
        return mime_type in self.mime_types or file_name.endswith(
            (".odt", ".ods", ".odp")
        )

    def extract(self, file: IO[bytes], max_characters: int) -> str:
        archive = _open_archive(file)
        if archive is None:
            return ""
        with archive, archive.open("content.xml") as content:
            return _join_until(self._lines(content), "\n", max_characters)

    def _lines(self, content: IO[bytes]) -> Iterable[str]:
        paragraph_tags = {
            f"{OPEN_DOCUMENT_TEXT_NAMESPACE}p",
            f"{OPEN_DOCUMENT_TEXT_NAMESPACE}h",
        }
        # Spaces, tabs and line breaks can be elements instead of characters.
        space_tags = {
            f"{OPEN_DOCUMENT_TEXT_NAMESPACE}s",
            f"{OPEN_DOCUMENT_TEXT_NAMESPACE}tab",
            f"{OPEN_DOCUMENT_TEXT_NAMESPACE}line-break",
        }
        row_tag = f"{OPEN_DOCUMENT_TABLE_NAMESPACE}table-row"
        cell_tag = f"{OPEN_DOCUMENT_TABLE_NAMESPACE}table-cell"
        # The text parts of the paragraphs that are being parsed, of which the last
        # one is nested in the others, like the paragraph of a frame.
        paragraphs = []
        # The cells of the current table row, which are combined into one line
        # like the rows of other spreadsheets.
        cell_texts = None
        cell_paragraphs = None
        # The text before the next element or after the previous one is only known
        # at the next event, because the elements are removed after they're parsed.
        pending_text = None

        for event, element in _iterparse_releasing(content):
            if pending_text is not None and paragraphs:
                paragraphs[-1].append(getattr(*pending_text) or "")
            if event == "start":
                pending_text = (element, "text")
                if element.tag in paragraph_tags:
                    paragraphs.append([])
                elif element.tag == row_tag:
                    cell_texts = []
                elif element.tag == cell_tag and cell_texts is not None:
                    cell_paragraphs = []
                continue

            pending_text = (element, "tail")
            if element.tag in space_tags:
                if paragraphs:
                    paragraphs[-1].append(" ")
            elif element.tag in paragraph_tags:
                text = "".join(paragraphs.pop())
                if cell_paragraphs is not None:
                    cell_paragraphs.append(text.strip())
                elif cell_texts is None and text.strip():
                    yield text
            elif element.tag == cell_tag and cell_paragraphs is not None:
                if text := " ".join(cell_paragraphs).strip():
                    cell_texts.append(text)
                cell_paragraphs = None
            elif element.tag == row_tag:
                if cell_texts:
                    yield "\t".join(cell_texts)
                cell_texts = None


class EmailFileTextExtractorType(FileTextExtractorType):
    type = "email"

    def can_extract(self, mime_type: str, file_name: str) -> bool:
        return mime_type == "message/rfc822" or file_name.endswith(".eml")

    def extract(self, file: IO[bytes], max_characters: int) -> str:
        message = BytesParser(policy=policy.default).parse(file)
        lines = [
            f"{header}: {message[header]}"
            for header in ("Subject", "From", "To")
            if message[header]
        ]
        body = message.get_body(preferencelist=("plain", "html"))
        if body is not None:
            content = body.get_content()
            if body.get_content_type() == "text/html":
                content = _html_to_text(content, max_characters)
            lines.extend(["", content.strip()])
        return "\n".join(lines)[:max_characters]


def _cell_to_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _row_to_text(values: Iterable) -> str | None:
    texts = [text for text in map(_cell_to_text, values) if text]
    return "\t".join(texts) if texts else None


class SpreadsheetFileTextExtractorType(FileTextExtractorType):
    type = "spreadsheet"

    mime_types = {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.template",
        "application/vnd.ms-excel.sheet.macroenabled.12",
        "application/vnd.ms-excel",
    }

    def can_extract(self, mime_type: str, file_name: str) -> bool:
        return mime_type in self.mime_types or file_name.endswith(
            (".xlsx", ".xlsm", ".xltx", ".xls")
        )

    def extract(self, file: IO[bytes], max_characters: int) -> str:
        is_legacy = file.read(len(LEGACY_OFFICE_FILE_SIGNATURE)) == (
            LEGACY_OFFICE_FILE_SIGNATURE
        )
        file.seek(0)
        lines = (
            self._legacy_workbook_lines(file)
            if is_legacy
            else self._workbook_lines(file)
        )
        return _join_until(lines, "\n", max_characters)

    def _workbook_lines(self, file: IO[bytes]) -> Iterable[str]:
        from openpyxl import load_workbook

        archive = _open_archive(file)
        if archive is None:
            return
        archive.close()
        file.seek(0)

        workbook = load_workbook(file, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                yield sheet.title
                for row in sheet.iter_rows(values_only=True):
                    if (text := _row_to_text(row)) is not None:
                        yield text
        finally:
            workbook.close()

    def _legacy_workbook_lines(self, file: IO[bytes]) -> Iterable[str]:
        from xlrd import open_workbook

        workbook = open_workbook(file_contents=file.read(), on_demand=True)
        try:
            for sheet in workbook.sheets():
                yield sheet.name
                for row in sheet.get_rows():
                    if (text := _row_to_text(cell.value for cell in row)) is not None:
                        yield text
        finally:
            workbook.release_resources()
